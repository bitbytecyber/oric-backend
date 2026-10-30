"""Import an HEC "ORIC Self-Assessment Scorecard" workbook (Index-A/B/C/D) into submissions.

The workbook is a stack of titled tables per sheet. For every table we
  1. find its title row "(vi) Research Proposals Approved …" and the following "SN" header row,
  2. map each header cell to a registry field via ``TABLES[...]["columns"]`` (regex on header text),
  3. read data rows until the next title, skipping instruction / grand-total rows and turning
     "National …" / "International …" sub-heading rows into the scope of the rows below,
  4. clean values (amounts, dates, nulls), parse the PI / inventor cell and match it to a Person,
  5. create one *verified* Submission per row (``source=excel_import``) keeping the complete
     original row under ``data["_source"]`` so no data point is lost.

Rows whose people cannot be matched are still created (they count on the scorecard) but get
``needs_match=True`` and an ``ImportRow(status=unmatched)`` with suggestions for the review queue.
"""
from __future__ import annotations

import datetime
import re
from dataclasses import dataclass, field

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from submissions.models import Contributor, ReviewEvent, Submission
from submissions.registry import CATEGORY_BY_KEY, lead_slot
from submissions.services import resolve_organization

from ..models import ImportBatch, ImportRow
from .normalize import (
    PersonMatcher,
    date_range,
    is_null,
    million,
    nat_int,
    rupees,
    split_people,
    text,
    to_date,
    to_number,
)

TITLE_RE = re.compile(r"^\s*\(([ivx]+(?:-[abc])?)\)\s*(.*)", re.I | re.S)

# Column target conventions:
#   field name              -> cleaned by the registry field type
#   "_pi" / "_lead"          -> lead contributor string (PI / inventor / winner)
#   "_copi"                  -> co-contributors string
#   "_duration"              -> "start to end" text -> start_date / end_date
#   "_evidence"              -> evidence link(s)
#   "_amount_rupees"         -> amount already in rupees (overrides the million column)
#   "_thematic_title"        -> "Thematic area and title" combined cell -> research_proposal_title
#   "_scope"                 -> National / International
#   "_date_duration"         -> "date and duration of agreement"
#   None                     -> kept only in data["_source"]
GRANT_COMMON = [
    (r"name of (joint )?research grant", "research_grant_name"),
    (r"name of co-?pi", "_copi"),
    (r"name (and details )?of pi|pi with designation", "_pi"),
    (r"thematic area and title", "_thematic_title"),
    (r"thematic area", "thematic_area"),
    (r"title of research proposal", "research_proposal_title"),
    (r"^duration", "_duration"),
    (r"collaborating partner", "collaborating_partners"),
    (r"co funding|co-funding", "co_funding_partners"),
    (r"national or international|national /\s*international|international or national", "_scope"),
    (r"annex", "_evidence"),
    (r"remarks", "remarks"),
    (r"key project deliverables", "key_project_deliverables"),
]

TABLES: dict[str, dict] = {
    # ---------------- Index-B -----------------------------------------------------------------
    "B.i": {"sheet": "Index-B", "label": "i", "category": "grant_opportunity", "columns": [
        (r"title of research grant opportunity", "opportunity_title"), (r"date of circulation", "circulation_date"),
        (r"national", "_scope"), (r"thematic", "thematic_area"), (r"title of research proposal",
                                                                  "research_proposal_title"),
        (r"starting date", "start_date"), (r"ending date", "end_date"),
        (r"collaborating", "collaborating_partners"), (r"co funding", "co_funding_partners"),
        (r"status", "award_status"), (r"remarks", "remarks"), (r"annex", "_evidence")]},
    "B.ii": {"sheet": "Index-B", "label": "ii", "category": "grant_hec_submitted", "columns": [
        (r"date of proposal submission", "proposal_submission_date"), (r"total funding requested",
                                                                       "total_funding_requested"),
        (r"^status", "status"), *GRANT_COMMON]},
    "B.iii": {"sheet": "Index-B", "label": "iii", "category": "grant_hec_approved", "columns": [
        (r"date of proposal approval", "proposal_approval_date"), (r"total funding approved",
                                                                   "total_funding_approved"),
        (r"oric overhead", "oric_overhead"), (r"date of approval", "?proposal_approval_date"), *GRANT_COMMON]},
    "B.iv": {"sheet": "Index-B", "label": "iv", "category": "grant_hec_completed", "columns": [
        (r"date of project completion", "project_completion_date"), (r"status of project", "project_status"),
        (r"total funding utilized", "total_funding_utilized"), (r"total funding released", "total_funding_released"),
        *GRANT_COMMON]},
    "B.v": {"sheet": "Index-B", "label": "v", "category": "grant_nonhec_submitted", "columns": [
        (r"date of proposal submission", "proposal_submission_date"), (r"total funding requested",
                                                                       "total_funding_requested"),
        (r"^status", "status"), *GRANT_COMMON]},
    "B.vi": {"sheet": "Index-B", "label": "vi", "category": "grant_nonhec_approved", "columns": [
        (r"date of proposal approval", "proposal_approval_date"), (r"total funding approved",
                                                                   "total_funding_approved"),
        (r"date of approval", "?proposal_approval_date"), *GRANT_COMMON]},
    "B.vii": {"sheet": "Index-B", "label": "vii", "category": "grant_nonhec_completed", "columns": [
        (r"name of research grant", "research_grant_name"),
        (r"date of project completion", "project_completion_date"),
        (r"total funding utilized", "total_funding_utilized"), (r"total funding released", "total_funding_released"),
        *GRANT_COMMON]},
    "B.viii-a": {"sheet": "Index-B", "label": "viii-a", "category": "joint_submitted", "columns": [
        (r"date of joint project submission", "submission_date"), (r"total funding requested",
                                                                   "total_funding_requested"),
        (r"status", "status"), *GRANT_COMMON]},
    "B.viii-b": {"sheet": "Index-B", "label": "viii-b", "category": "joint_approved", "columns": [
        (r"date of joint project approval", "approval_date"), (r"total funding approved", "total_funding_approved"),
        (r"status", "status"), *GRANT_COMMON]},
    "B.viii-c": {"sheet": "Index-B", "label": "viii-c", "category": "joint_completed", "columns": [
        (r"date of joint project completion", "completion_date"),
        (r"total funding utilized", "total_funding_utilized"), (r"total funding released", "total_funding_released"),
        (r"status of project", "project_status"), *GRANT_COMMON]},
    "B.x": {"sheet": "Index-B", "label": "x", "category": "irb_review", "columns": [
        (r"date of review", "review_date"), (r"total amount requested", "total_amount_requested"),
        (r"irb meeting decision", "irb_decision"), (r"key project expected", "key_deliverables"), *GRANT_COMMON]},
    "B.xi": {"sheet": "Index-B", "label": "xi", "category": "policy_advocacy", "columns": [
        (r"name of government body", "government_body_presented"), (r"date of presentation", "presentation_date"),
        (r"area advocated", "advocacy_area"), (r"^brief", "brief"), (r"coalition", "coalition_partners"),
        (r"issue verification", "research_status"), (r"advocacy tools", "advocacy_tools"),
        (r"name of pi", "_pi"), (r"^duration", "_duration"), (r"annex", "_evidence")]},
    "B.xii": {"sheet": "Index-B", "label": "xii", "category": "research_link", "columns": [
        (r"type of linkage", "linkage_type"), (r"date of mou", "mou_agreement_date"),
        (r"international or national|national", "_scope"),
        (r"host institution", "host_institution"), (r"collaborating agency", "collaborating_agency"),
        (r"field of study", "field_of_study"), (r"scope of collaboration", "scope_of_collaboration"),
        (r"salient features", "salient_features"), (r"annex", "_evidence")]},
    "B.xiii": {"sheet": "Index-B", "label": "xiii", "category": "contract_research", "scope": "National",
               "columns": [
                   (r"thematic areas and title", "_thematic_title"), (r"date of contract signed",
                                                                       "contract_signed_date"),
                   (r"name of co-pi", "_copi"), (r"name of pi", "_pi"), (r"sponsoring agency", "sponsoring_agency"),
                   (r"national", "_scope"), (r"counterpart from industry", "counterpart_industry"),
                   (r"^duration", "_duration"), (r"total amount approved", "total_amount_approved"),
                   (r"^__col_L$", "_amount_rupees"), (r"expected deliverables", "expected_deliverables"),
                   (r"date of contract$", "?contract_signed_date"), (r"annex", "_evidence")]},
    "B.xvi": {"sheet": "Index-B", "label": "xvi", "category": "civic_engagement", "columns": [
        (r"title of event", "event_title"), (r"date of event", "event_date"),
        (r"component of community", "community_component"), (r"^outcome", "outcome"),
        (r"collaboration developed", "collaboration_developed"), (r"csos", "engaged_csos_ngos"),
        (r"sponsoring agency", "sponsoring_agency"), (r"arranged", "arranged_or_participated"),
        (r"dissemination", "dissemination_material"), (r"remarks", "remarks"), (r"annex", "_evidence")]},
    "B.xvii": {"sheet": "Index-B", "label": "xvii", "category": "consultancy_opportunity", "columns": [
        (r"title of project", "project_title"), (r"date of circulation|date of execution", "circulation_date"),
        (r"name of pi", "_pi"), (r"company details", "company_details"),
        (r"contract value", "contract_value"), (r"project timelines", "_duration"),
        (r"type of consultancy", "consultancy_type"), (r"key deliverables", "key_deliverables"),
        (r"oric percentage", None), (r"remarks", "remarks"), (r"annex", "_evidence")]},
    "B.xviii": {"sheet": "Index-B", "label": "xviii", "category": "consultancy_contract", "columns": [
        (r"title of project", "project_title"), (r"date of execution", "execution_date"),
        (r"name of pi", "_pi"), (r"company details", "company_details"),
        (r"contract value \(pkr millions\)", "contract_value"), (r"contract value \(pkr\)", "_amount_rupees"),
        (r"project timelines", "_duration"), (r"type of consultancy", "consultancy_type"),
        (r"key deliverables", "key_deliverables"), (r"oric percentage", "oric_percentage"),
        (r"remarks", "remarks"), (r"annex", "_evidence")]},
    "B.xix": {"sheet": "Index-B", "label": "xix", "category": "asrb_liaison", "columns": [
        (r"liaison developed with", "liaison_with"), (r"date of execution", "execution_date"),
        (r"annex", "_evidence")]},
    # ---------------- Index-C -----------------------------------------------------------------
    "C.i": {"sheet": "Index-C", "label": "i", "category": "policy", "fixed": {"policy_type": "Research policy"},
            "columns": [(r"title of", "policy_title"), (r"ref", "policy_ref_no"), (r"date of approval",
                                                                                   "approval_date"),
                        (r"approved by", "approved_by")]},
    "C.ii": {"sheet": "Index-C", "label": "ii", "category": "policy", "fixed": {"policy_type": "IP policy"},
             "columns": [(r"title of", "policy_title"), (r"ref", "policy_ref_no"), (r"date of approval",
                                                                                    "approval_date"),
                         (r"approved by", "approved_by")]},
    "C.iii": {"sheet": "Index-C", "label": "iii", "category": "policy", "fixed": {"policy_type": "Other policy"},
              "columns": [(r"title of", "policy_title"), (r"ref", "policy_ref_no"), (r"date of approval",
                                                                                     "approval_date"),
                          (r"approved by", "approved_by")]},
}

INVENTION = [
    (r"lead inventor", "_lead"), (r"title of invention", "invention_title"), (r"category of ip", "ip_category"),
    (r"development status", "development_status"), (r"key scientific", "key_scientific_aspects"),
    (r"commercial partner", "commercial_partner"), (r"financial support", "financial_support"),
    (r"national|international", "_scope"), (r"annex|^__col_L$", "_evidence"),
]
TABLES.update({
    "C.iv": {"sheet": "Index-C", "label": "iv", "category": "ip_disclosure", "columns": [
        (r"disclosure made with", "disclosure_made_with"), (r"previous disclosure", "previous_disclosure"),
        *INVENTION]},
    **{f"C.{lbl}": {"sheet": "Index-C", "label": lbl, "category": "patent",
                    "fixed": {"filed_or_granted": fg}, "scope": sc, "columns": [
                        (r"patent filed with|granting authority", "filed_with"), (r"date of filing", "filing_date"),
                        *INVENTION]}
       for lbl, fg, sc in [("v", "Filed", "National"), ("vi", "Granted", "National"),
                           ("vii", "Filed", "International"), ("viii", "Granted", "International")]},
    **{f"C.{lbl}": {"sheet": "Index-C", "label": lbl, "category": "licensing_negotiation", "scope": sc, "columns": [
        (r"field of use", "field_of_use"), (r"duration of agreement", "duration_of_agreement"),
        (r"licensee", "licensee"), (r"status of negotiation", "negotiation_status"), *INVENTION]}
       for lbl, sc in [("ix", "National"), ("xi", "International")]},
    **{f"C.{lbl}": {"sheet": "Index-C", "label": lbl, "category": "license_signed", "scope": sc, "columns": [
        (r"field of use", "field_of_use"), (r"date and duration", "_date_duration"), (r"licensee", "licensee"),
        (r"exclusive", "license_type"), *INVENTION]}
       for lbl, sc in [("x", "National"), ("xii", "International")]},
    "C.xiii": {"sheet": "Index-C", "label": "xiii", "category": "prototype", "columns": [
        (r"field of use", "field_of_use"), (r"industrial partner", "industrial_partner"), (r"remarks", "remarks"),
        *INVENTION]},
    "C.xiv": {"sheet": "Index-C", "label": "xiv", "category": "creative_product", "columns": [
        (r"^title", "title"), (r"department of lead|name.*lead", "_lead"), (r"^category", "product_category"),
        (r"national", "_scope"), (r"forum", "forum"), (r"^status", "status"),
        (r"financial support", "financial_support"), (r"field of use", "field_of_use"), (r"annex", "_evidence")]},
    "C.xv": {"sheet": "Index-C", "label": "xv", "category": "industry_visit", "columns": [
        (r"name of visitor", "visitor_name"), (r"date of visit", "visit_date"), (r"agenda", "agenda"),
        (r"annex|^__col_L$", "_evidence")]},
    "C.xvi": {"sheet": "Index-C", "label": "xvi", "category": "collaboration_agreement", "columns": [
        (r"type of linkage", "linkage_type"), (r"international or national|national", "_scope"),
        (r"host institution", "partner_institution"), (r"^duration", "duration"),
        (r"key initiatives", "key_initiatives"), (r"^field", "field"), (r"scope of collaboration",
                                                                        "scope_of_collaboration"),
        (r"establishment date", "establishment_date"), (r"financial support", "financial_support"),
        (r"annex", "_evidence")]},
    "C.xvii": {"sheet": "Index-C", "label": "xvii", "category": "award", "columns": [
        (r"title of award", "title"), (r"forum", "conferring_organization"), (r"award/ prize|award/prize|any award",
                                                                               "award_received"),
        (r"brief details", "work_details"), (r"prize money", "prize_money"), (r"award winner", "_lead"),
        (r"remarks", "remarks"), (r"annex", "_evidence")]},
    # ---------------- Index-D -----------------------------------------------------------------
    "D.i": {"sheet": "Index-D", "label": "i", "category": "research_revenue", "columns": [
        (r"name of research grant", "research_grant_name"), (r"national or international", "_scope_or_date"),
        (r"title of research proposal", "research_proposal_title"), (r"name and details of pi", "_pi"),
        (r"^duration", "_duration"), (r"total funding approved", "total_funding_approved"),
        (r"total funding released", "total_funding_released"),
        (r"overhead in approved", "oric_overhead_approved"), (r"overhead in released", "oric_overhead_released"),
        (r"remarks", "remarks"), (r"annex", "_evidence")]},
    "D.ii": {"sheet": "Index-D", "label": "ii", "category": "licensing_revenue", "columns": [
        (r"lead inventor", "_lead"), (r"title of invention", "invention_title"), (r"national", "_scope"),
        (r"date and duration", "_date_duration"), (r"licensee", "licensee"), (r"share percentages",
                                                                               "share_percentages"),
        (r"total amount", "total_amount"), (r"oric approved share", "oric_approved_share"),
        (r"oric received share", "oric_received_share"), (r"remarks", "remarks"), (r"annex", "_evidence")]},
    **{f"D.{lbl}": {"sheet": "Index-D", "label": lbl, "category": cat, "subheadings": True, "fixed": fixed,
                    "columns": [
                        (r"title of training|type of event", "title"), (r"date of event", "event_date"),
                        (r"no\. of participants", "participants"), (r"focus area", "focus_and_outcomes"),
                        (r"^organizer", "organizer"), (r"audience", "audience_type"),
                        (r"oric personnel", "oric_personnel"), (r"annex", "_evidence")]}
       for lbl, cat, fixed in [("iii", "training_oric", {"event_type": "Training"}),
                               ("iv", "training_external", {"event_type": "Training"}),
                               ("v", "exhibition", {"event_type": "Exhibition"}),
                               ("vi", "training_oric_staff", {"event_type": "Training"})]},
})

TABLES["B.xiv"] = {**TABLES["B.xiii"], "label": "xiv", "scope": "International"}

SKIP_TABLES = {"B.ix": "computed from B.vi (R&D funding volume)", "B.xv": "computed from B.xiii + B.xiv"}

# ---------------------------------------------------------------------------
# Templates. HEC changed the workbook three times:
#   annexure_2019 – FY 2019-20 "Performance Matrices": Score Card + Annexure 1-4, own numbering
#   index_2020    – FY 2020-21 Index-A..D, but different item numbering (and "Index D" with a space)
#   index_2021    – FY 2021-22 onward, the layout ``TABLES`` was written for
# Tables are therefore recognised by their title text, not by the roman numeral in front of it.
# ---------------------------------------------------------------------------
TEMPLATE_FRAMEWORK = {"annexure_2019": "hec_2019", "index_2020": "hec_2020", "index_2021": "hec_2021"}

# (index letter, title regex, canonical key) – first match wins, so specific patterns come first.
TITLE_KEYS = [
    ("B", r"grant opportunities", "B.i"),
    ("B", r"submitted for funding (by|from) hec", "B.ii"),
    ("B", r"approved for funding (by|from) hec", "B.iii"),
    ("B", r"\(won through hec\)", "B.iv"),
    ("B", r"submitted for funding \(from non.?hec", "B.v"),
    ("B", r"approved for funding \(from non.?hec", "B.vi"),
    ("B", r"\(won through non.?hec", "B.vii"),
    ("B", r"^joint research projects submitted", "B.viii-a"),
    ("B", r"^joint research projects approved", "B.viii-b"),
    ("B", r"^joint research projects completed", "B.viii-c"),
    ("B", r"volume of r&d funding", "B.ix"),
    ("B", r"reviewed by irb", "B.x"),
    ("B", r"policy advocacy", "B.xi"),
    ("B", r"research links established", "B.xii"),
    ("B", r"volume of contract research", "B.xv"),
    ("B", r"contract research awarded.*\(international\)", "B.xiv"),
    ("B", r"contract research awarded", "B.xiii"),
    ("B", r"civic engagement", "B.xvi"),
    ("B", r"consultancy opportunities", "B.xvii"),
    ("B", r"consultancy contracts", "B.xviii"),
    ("B", r"liaison developed", "B.xix"),
    ("C", r"^research policy", "C.i"),
    ("C", r"intellectual property rights policy", "C.ii"),
    ("C", r"research ethics policy", "C.iii"),
    ("C", r"ip disclosures", "C.iv"),
    ("C", r"filed \(at international", "C.vii"),
    ("C", r"granted \(at international", "C.viii"),
    ("C", r"\bfiled\b", "C.v"),
    ("C", r"\bgranted\b", "C.vi"),
    ("C", r"negotiations initiated \(at international", "C.xi"),
    ("C", r"licensing negotiations", "C.ix"),
    ("C", r"licenses signed \(at internat", "C.xii"),
    ("C", r"licenses signed", "C.x"),
    ("C", r"research products", "C.xiii"),
    ("C", r"science / arts", "C.xiv"),
    ("C", r"visits by representatives", "C.xv"),
    ("C", r"agreements signed for collaboration", "C.xvi"),
    ("C", r"hono(u)?rs or awards", "C.xvii"),
    ("D", r"annual research revenue", "D.i"),
    ("D", r"revenue earned", "D.ii"),
    ("D", r"^exhibitions", "D.v"),  # its title ends "… Arranged by ORIC", so it goes before D.iii
    ("D", r"arranged by oric", "D.iii"),
    ("D", r"arranged by other heis", "D.iv"),
    ("D", r"arranged on research", "D.vi"),
]

# FY 2019-20 annexure tables: (sheet, title regex, key)
ANNEX_TITLE_KEYS = [
    ("Annexure 2", r"approved for funding by hec", "A2.i"),
    ("Annexure 2", r"approved for funding \(other than hec", "A2.ii"),
    ("Annexure 2", r"joint research projects approved", "A2.iii"),
    ("Annexure 2", r"projects completed", "A2.iv"),
    ("Annexure 2", r"r&d funding secured", "A2.v"),
    ("Annexure 2", r"research links established", "A2.vi"),
    ("Annexure 2", r"contract research awarded", "A2.vii"),
    ("Annexure 2", r"research database", "A2.viii"),
    ("Annexure 2", r"advisory board", "A2.ix"),
    ("Annexure 3", r"trainings/ ?workshops", "A3.i"),
    ("Annexure 3", r"ideas competition", "A3.ii"),
    ("Annexure 3", r"participation of students", "A3.iii"),
    ("Annexure 3", r"participation of faculty", "A3.iv"),
    ("Annexure 3", r"attended by oric personnel", "A3.v"),
    ("Annexure 3", r"incorporated entrepreneurship", "A3.vi"),
    ("Annexure 3", r"revenue generated from trainings", "A3.vii"),
    ("Annexure 4", r"development of research", "A4.i"),
    ("Annexure 4", r"revenue generated from commerciali", "A4.ii"),
    ("Annexure 4", r"\bfiled\b", "A4.iii"),
    ("Annexure 4", r"\bgranted\b", "A4.iv"),
    ("Annexure 4", r"research products", "A4.v"),
    ("Annexure 4", r"patent licensed", "A4.vi"),
    ("Annexure 4", r"incubation cent", "A4.vii"),
    ("Annexure 4", r"links established with corporate", "A4.viii"),
]

HEC_GRANT_RE = re.compile(r"\b(nrpu|ttsf|gcf|icrg|lcf|tdf|srgp|hec)\b", re.I)


def _policy_type(data: dict) -> dict:
    t = (data.get("policy_title") or "").lower()
    if "intellectual property" in t or re.search(r"\bip\b", t):
        return {"policy_type": "IP policy"}
    if "research" in t:
        return {"policy_type": "Research policy"}
    return {"policy_type": "Other policy"}


GRANT_2019 = [
    (r"name of co-?pi", "_copi"), (r"name of pi", "_pi"), (r"^department", "_dept"), (r"^designation", "_desig"),
    (r"field of study", "thematic_area"), (r"title of research proposal", "research_proposal_title"),
    (r"^duration", "_duration"), (r"co-funding agency", "co_funding_partners"), (r"co-funding details", None),
    (r"installment", None), (r"ref\. ?no|reference", None),
]
ANNEX_TABLES: dict[str, dict] = {
    "A2.i": {"category": "grant_hec_approved", "columns": [
        (r"title of project", "research_grant_name"), (r"total amount approved", "total_funding_approved"),
        (r"oric overhead", "oric_overhead"), (r"date of approval", "proposal_approval_date"), *GRANT_2019]},
    "A2.ii": {"category": "grant_nonhec_approved", "columns": [
        (r"title of project", "research_grant_name"), (r"name of funding agency", "funding_agency"),
        (r"total amount approved", "total_funding_approved"), (r"date of approval", "proposal_approval_date"),
        *GRANT_2019]},
    "A2.iii": {"category": "joint_approved", "columns": [
        (r"title of project", "research_grant_name"), (r"total amount approved", "total_funding_approved"),
        (r"details of share", None), (r"date of approval", "approval_date"), (r"sponsoring agency", "funding_agency"),
        (r"collaborating partners", "collaborating_partners"), (r"approval letter", None), *GRANT_2019]},
    "A2.iv": {"category": "grant_nonhec_completed",
              "category_fn": lambda d: "grant_hec_completed" if HEC_GRANT_RE.search(d.get("research_grant_name") or "")
              else "grant_nonhec_completed",
              "columns": [
                  (r"title of project", "research_grant_name"), (r"total amount approved", None),
                  (r"oric overhead", None), (r"project completion date", "project_completion_date"),
                  (r"co-funding details", "co_funding_partners"), (r"deliverables", "key_project_deliverables"),
                  (r"audited", None), *GRANT_2019]},
    "A2.vi": {"category": "research_link", "columns": [
        (r"^0$|type of linkage", "linkage_type"), (r"international or national", "_scope"),
        (r"host institution", "host_institution"), (r"collaborating", "collaborating_agency"),
        (r"field of study", "field_of_study"), (r"scope of collaboration", "scope_of_collaboration"),
        (r"linkage establishment date", "mou_agreement_date"), (r"soft copy", "_evidence")]},
    "A2.vii": {"category": "contract_research", "columns": [
        (r"title of research proposal", "research_proposal_title"), (r"purpose of research", None),
        (r"name of co-?pi", "_copi"), (r"name of pi", "_pi"), (r"field of study", "thematic_area"),
        (r"counterpart from industry", "counterpart_industry"), (r"^duration", "_duration"),
        (r"total amount approved", "total_amount_approved"), (r"expected deliverables", "expected_deliverables"),
        (r"details of share", None), (r"date of agreement", "contract_signed_date"),
        (r"sponsoring agency", "sponsoring_agency"), (r"agreement reference", "_evidence")]},
    "A2.ix": {"category": "advisory_board_meeting", "require_sn": True, "columns": [
        (r"date of me", "meeting_date"), (r"minutes", "minutes_date"), (r"actions suggested", "actions_suggested"),
        (r"action compliance", "action_compliance"), (r"list of cia|board members", "members")]},
    "A3.i": {"category": "training_oric", "subheadings": True, "fixed": {"event_type": "Training"}, "columns": [
        (r"title of training", "title"), (r"date of event", "event_date"), (r"no\. of participants", "participants"),
        (r"focus area", "focus_and_outcomes"), (r"audience", "audience_type")]},
    "A3.ii": {"category": "innovation_competition", "columns": [
        (r"title of competition", "title"), (r"date of event", "event_date"),
        (r"no\. of participants", "participants"), (r"focus area", "focus_and_outcomes")]},
    "A3.iii": {"category": "training_external", "fixed": {"event_type": "Training"}, "columns": [
        (r"title of the training", "title"), (r"date of event", "event_date"), (r"name of hei", "organizer"),
        (r"number of participants", "participants"), (r"audience type", "audience_type"),
        (r"major focus area", "focus_and_outcomes"), (r"major area of training", None)]},
    "A3.iv": {"category": "training_external", "fixed": {"audience_type": "Faculty"}, "columns": [
        (r"international or national", "_scope"), (r"^type of event", "event_type"), (r"title of the event", "title"),
        (r"event date", "event_date"), (r"event venue", None), (r"event organizer", "_lead"),
        (r"field of study", None), (r"outcome", "focus_and_outcomes"), (r"sanctioned amount", None),
        (r"funding agencies", None), (r"soft copy", "_evidence")]},
    "A3.v": {"category": "training_oric_staff", "columns": [
        (r"title of workshop", "title"), (r"type of training", "event_type"), (r"starting date", "event_date"),
        (r"ending date", None), (r"field of study", "focus_and_outcomes"), (r"participants type", "oric_personnel"),
        (r"expenditure", None)]},
    "A3.vi": {"category": "entrepreneurship_course", "columns": [
        (r"title of course", "course_title"), (r"course level", "course_level"), (r"credit hours", "credit_hours"),
        (r"number of disciplines", "disciplines")]},
    "A4.i": {"category": "policy", "fixed_fn": _policy_type, "columns": [
        (r"^title of", "policy_title"), (r"draft status", None), (r"date of approval", "approval_date"),
        (r"policy addressing", None)]},
    "A4.iii": {"category": "patent", "fixed": {"filed_or_granted": "Filed"}, "scope_from": "filed_with", "columns": [
        (r"title of patent", "invention_title"), (r"name of pi", "_lead"), (r"research field", "key_scientific_aspects"),
        (r"issuance authority", "filed_with"), (r"approving authority", None), (r"date of filing", "filing_date"),
        (r"documentary", "_evidence")]},
    "A4.iv": {"category": "patent", "fixed": {"filed_or_granted": "Granted"}, "scope_from": "filed_with",
              "columns": [
        (r"title of patent", "invention_title"), (r"name of pi", "_lead"), (r"research field", "key_scientific_aspects"),
        (r"issuance authority", "filed_with"), (r"approving authority", None), (r"date of filing", "filing_date"),
        (r"documentary", "application_no")]},
    "A4.v": {"category": "prototype", "columns": [
        (r"title / name", "invention_title"), (r"name and designation of pi", "_lead"),
        (r"research field", "field_of_use"), (r"research project title", "key_scientific_aspects"),
        (r"collaborating indust", "industrial_partner"), (r"funding agency", "financial_support"),
        (r"current status", "development_status")]},
    "A4.vi": {"category": "license_signed", "columns": [
        (r"title of patent", "invention_title"), (r"name and designation of pi", "_lead"),
        (r"research field", "field_of_use"), (r"industry name", "licensee"), (r"date of issuance", "agreement_date"),
        (r"patent details", None), (r"documentary", "_evidence")]},
    "A4.vii": {"category": "incubation_center", "aggregate": True, "columns": [
        (r"status of establishment", "status"), (r"companies incubated", "companies_incubated"),
        (r"companies graduated", "companies_graduated"), (r"revenue generated", "revenue_details"),
        (r"business plans", "business_plans"), (r"details of ic hr", "hr_details")]},
    "A4.viii": {"category": "collaboration_agreement", "columns": [
        (r"title / name of the corporate", "partner_institution"), (r"research field", "field"),
        (r"date of signing", "establishment_date"), (r"documentary", "_evidence")]},
}
ANNEX_SKIP = {"A2.v": "computed from A2.i-iii (R&D funding volume)", "A2.viii": "ORIC profile (database), not a list",
              "A3.vii": "no portal category (training revenue), table is empty",
              "A4.ii": "no portal category (commercialisation revenue), table is empty"}

# Sheets that are scratch copies / alternates and never imported unless named in ``sheet_overrides``.
IGNORED_SHEETS = re.compile(r"^(sheet\d*|extra|copy of .*|uploade? pc)$", re.I)
IP_FAMILY = ("ip_disclosure", "patent", "licensing_negotiation", "license_signed", "prototype")


def canon_sheet(name: str) -> str:
    """'Index D' / 'index-d ' -> 'Index-D'; other names unchanged."""
    m = re.match(r"^\s*index\s*-?\s*([a-d])\s*$", name, re.I)
    return f"Index-{m.group(1).upper()}" if m else name.strip()


def _title_key(prefix: str, title: str) -> str | None:
    t = " ".join(title.split()).lower()
    return next((key for p, pattern, key in TITLE_KEYS if p == prefix and re.search(pattern, t)), None)


def detect_template(wb) -> str:
    names = {canon_sheet(n) for n in wb.sheetnames}
    if "Index-B" not in names and any(n.lower().startswith("annexure") for n in names):
        return "annexure_2019"
    if "Index-B" not in names:
        raise ValueError("Not an HEC ORIC scorecard workbook: no Index-B sheet and no Annexure sheets")
    ws = wb["Index-B"]
    for r in range(1, min(ws.max_row, 12) + 1):
        for c in (1, 2, 3):
            v = ws.cell(r, c).value
            if isinstance(v, str) and re.search(r"grant opportunities", v, re.I):
                return "index_2021"
    return "index_2020"


class _Cell:
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


_EMPTY = _Cell(None)


class Grid:
    """A worksheet loaded once into memory (values only, at most ``max_cols`` columns).

    Several workbooks declare ~1M rows / 16k columns of formatting, which makes openpyxl's cell access crawl.
    This keeps the ``ws.cell(r, c).value`` call style with the real data bounds."""

    def __init__(self, ws, max_cols=40, max_blank_run=500):
        rows, blank = [], 0
        for row in ws.iter_rows(values_only=True, max_col=max_cols):
            row = list(row)
            if all(v in (None, "") for v in row):
                blank += 1
                if blank > max_blank_run:
                    break
            else:
                blank = 0
            rows.append(row)
        while rows and all(v in (None, "") for v in rows[-1]):
            rows.pop()
        width = max((max((i + 1 for i, v in enumerate(r) if v not in (None, "")), default=0) for r in rows), default=0)
        self.rows = [r[:width] + [None] * (width - len(r[:width])) for r in rows]
        self.max_row, self.max_column = len(self.rows), width
        self.title = ws.title

    def cell(self, r, c):
        if 1 <= r <= self.max_row and 1 <= c <= self.max_column:
            return _Cell(self.rows[r - 1][c - 1])
        return _EMPTY


class Book:
    """Workbook with canonical sheet names ('Index D' -> 'Index-D') and lazily loaded grids."""

    def __init__(self, path):
        # not read_only: the streaming reader silently drops some cells in these files (FY 2024-25 Index-B!O680)
        self.wb = load_workbook(path, data_only=True)
        self.names = {canon_sheet(n): n for n in self.wb.sheetnames}
        self.sheetnames = list(self.names)
        self._grids: dict[str, Grid] = {}

    def __contains__(self, name):
        return name in self.names

    def __getitem__(self, name) -> Grid:
        if name not in self._grids:
            self._grids[name] = Grid(self.wb[self.names[name]])
        return self._grids[name]


@dataclass
class TableHit:
    key: str
    sheet: str
    title_row: int
    header_row: int
    end_row: int
    headers: dict = field(default_factory=dict)  # column letter -> header text
    label: str = ""  # roman numeral as printed in the workbook
    title: str = ""


def _annex_key(sheet: str, title: str) -> str | None:
    t = " ".join(title.split()).lower()
    return next((key for s, pattern, key in ANNEX_TITLE_KEYS if s == sheet and re.search(pattern, t)), None)


def locate_tables(wb, template: str = "index_2021", sheet_overrides: dict | None = None) -> list[TableHit]:
    """Find every titled table. ``sheet_overrides`` = {"B.xi": "Uploade PC"} reads that table from another sheet."""
    overrides = {k: canon_sheet(v) for k, v in (sheet_overrides or {}).items()}
    if template == "annexure_2019":
        sheets = [(n, None) for n in wb.sheetnames if n.lower().startswith("annexure") and n != "Annexure 1"]
    else:
        sheets = [(n, n[-1]) for n in ("Index-B", "Index-C", "Index-D") if n in wb]
    base = {n for n, _ in sheets}
    for key, sheet in overrides.items():
        if sheet not in wb:
            raise ValueError(f"Sheet {sheet!r} (override for {key}) is not in the workbook")
        if sheet not in base and all(n != sheet for n, _ in sheets):
            sheets.append((sheet, None if template == "annexure_2019" else key[0]))
    hits: list[TableHit] = []
    for sheet, prefix in sheets:
        ws = wb[sheet]
        titles = []
        for r in range(1, ws.max_row + 1):
            for c in (1, 2, 3):
                v = ws.cell(r, c).value
                if isinstance(v, str):
                    m = TITLE_RE.match(v)
                    if m:
                        titles.append((r, m.group(1).lower(), " ".join(m.group(2).split())))
                    break
                if v not in (None, ""):
                    break
        for i, (row, label, title) in enumerate(titles):
            if prefix is None:
                # an extra override sheet keeps its annexure title, so match the patterns of the key it stands in for
                key = _annex_key(sheet, title) or next(
                    (k for s_, pattern, k in ANNEX_TITLE_KEYS if overrides.get(k) == sheet
                     and re.search(pattern, " ".join(title.split()).lower())), None) or f"{sheet}.{label}"
            else:
                key = _title_key(prefix, title) or (f"{prefix}.{label}" if template == "index_2021" else
                                                    f"{prefix}.?{label}")
            if sheet not in base and overrides.get(key) != sheet:
                continue  # an extra override sheet contributes only the table it was named for
            if key in overrides and overrides[key] != sheet:
                continue  # replaced by the override sheet
            nxt = titles[i + 1][0] if i + 1 < len(titles) else ws.max_row + 1
            header_row = None
            for r in range(row + 1, min(row + 6, nxt)):
                if any(isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.strip().upper() in
                       ("SN", "S.N", "S.NO", "S. NO", "SR. NO.") for c in (1, 2)):
                    header_row = r
                    break
            if header_row is None:
                continue
            headers = {}
            for c in range(1, ws.max_column + 1):
                v = ws.cell(header_row, c).value
                headers[get_column_letter(c)] = " ".join(str(v).split()) if v not in (None, "") else ""
            hits.append(TableHit(key, sheet, row, header_row, nxt - 1, headers, label, title))
    return hits


def _column_map(spec_cols, headers: dict) -> dict:
    """Header -> field. A ``"?field"`` target is a fallback: used only when no other column maps to ``field``
    (FY 2020-21 has just "Date of Approval" where later files also carry "Date of Proposal Approval")."""
    out = {}
    used_targets = set()
    for col, head in headers.items():
        h = head.lower() if head else f"__col_{col}".lower()
        if h in ("sn", "s.n", "s.no", "sr. no.") or (not head and False):
            continue
        for pattern, target in spec_cols:
            pat = pattern.lower()
            if re.search(pat, h if head else f"__col_{col}".lower()):
                if target and target in used_targets and not target.startswith("_"):
                    continue
                out[col] = target
                if target and not target.startswith("?"):
                    used_targets.add(target)
                break
    for col, target in out.items():
        if target and target.startswith("?"):
            out[col] = None if target[1:] in used_targets else target[1:]
            used_targets.add(target[1:])
    return out


def _is_noise(values: list) -> bool:
    joined = " ".join(str(v) for v in values[:6] if v not in (None, "")).lower()
    if not joined.strip():
        return True
    return bool(re.search(r"grand total|^total\b|please add|add as many rows|sub total", joined))


def _subheading(values: list) -> str | None:
    filled = [v for v in values if v not in (None, "")]
    if 1 <= len(filled) <= 2 and isinstance(filled[0], str) and len(filled[0]) < 60:
        s = filled[0].lower()
        if "international" in s:
            return "International"
        if "national" in s:
            return "National"
    return None


@dataclass
class ImportResult:
    created: int = 0
    unmatched: int = 0
    skipped: int = 0
    errors: int = 0
    per_table: dict = field(default_factory=dict)
    flags: dict = field(default_factory=dict)  # table key -> number of rows with data-quality notes
    log: list = field(default_factory=list)


def import_workbook(batch: ImportBatch, path: str, *, replace: bool = True, user=None, only=None,
                    sheet_overrides: dict | None = None) -> dict:
    """Import one workbook into ``batch.cycle``.

    ``only`` limits the run to some table keys (their earlier rows in this year are replaced, everything else is
    kept); ``sheet_overrides`` reads a table from another sheet, e.g. {"B.xi": "Uploade PC"}."""
    wb = Book(path)
    template = detect_template(wb)
    cycle = batch.cycle
    matcher = PersonMatcher()
    res = ImportResult()
    only = set(only or [])
    specs, skips = (ANNEX_TABLES, ANNEX_SKIP) if template == "annexure_2019" else (TABLES, SKIP_TABLES)
    res.log.append(f"Template: {template} ({TEMPLATE_FRAMEWORK[template]})")
    as_submitted = score_card_as_submitted(wb, template)
    hits = locate_tables(wb, template, sheet_overrides)
    if only:
        # refuse before deleting anything: a typo or a table this template lacks would wipe it and import nothing
        found = {h.key for h in hits}
        bad = sorted(k for k in only if k not in found or k in skips or k not in specs)
        if bad:
            raise ValueError(f"Not importable from this {template} workbook: {', '.join(bad)}")
    with transaction.atomic():
        if replace:
            rows = ImportRow.objects.filter(batch__cycle=cycle, batch__kind=ImportBatch.Kind.HEC_WORKBOOK).exclude(
                batch=batch)
            subs = Submission.objects.filter(cycle=cycle, source=Submission.Source.EXCEL)
            if only:
                rows = rows.filter(table__in=only)
                # ids first: deleting the rows below would otherwise empty this join (ImportRow.submission is SET_NULL)
                subs = Submission.objects.filter(
                    pk__in=list(subs.filter(import_rows__table__in=only).values_list("pk", flat=True).distinct()))
            # earlier row logs (incl. unmatched-queue entries) go with the records they describe
            rows.delete()
            subs.delete()
        if not only:
            _import_index_a(wb, cycle, user, res, template=template, as_submitted=as_submitted)
        cycle.framework = TEMPLATE_FRAMEWORK[template]
        cycle.save(update_fields=["framework"])
        seen = set()
        for hit in hits:
            if only and hit.key not in only:
                continue
            if hit.key in skips:
                res.log.append(f"{hit.key}: skipped ({skips[hit.key]})")
                continue
            spec = specs.get(hit.key)
            if not spec or hit.key in seen:
                why = "duplicate table" if spec else "no mapping"
                res.log.append(f"{hit.key} ({hit.sheet} row {hit.title_row}, '{hit.title[:60]}'): {why}, skipped")
                continue
            seen.add(hit.key)
            if template == "index_2021" and hit.label and hit.key.split(".", 1)[1] != hit.label:
                res.log.append(f"{hit.key}: numbered ({hit.label}) in {hit.sheet}, recognised by its title")
            if spec.get("aggregate"):
                _import_aggregate(batch, wb[hit.sheet], hit, spec, res, user)
            else:
                _import_table(batch, wb[hit.sheet], hit, spec, matcher, res, user)
        if not only:
            _import_strategic_plan(batch, wb, cycle, user, res)
            if template == "index_2021":
                _sheet_scores(wb, res)
    stats = {"created": res.created, "unmatched": res.unmatched, "skipped": res.skipped, "errors": res.errors,
             "per_table": res.per_table, "sheet_scorecard": res.per_table.pop("_sheet_scorecard", {}),
             "template": template, "flags": res.flags, "as_submitted": as_submitted}
    if only:
        stats["only"] = sorted(only)
    if sheet_overrides:
        stats["sheet_overrides"] = sheet_overrides
    batch.stats = stats
    batch.log = "\n".join(res.log)
    batch.status = ImportBatch.Status.APPLIED
    batch.save()
    return stats


def _sane_date(d: str | None, name: str, raw, flags: list) -> str | None:
    if d and not (1990 <= int(d[:4]) <= 2100):
        flags.append(f"{name}: impossible date '{text(raw)}' dropped")
        return None
    return d


def _clean_date(v, name: str, flags: list) -> str | None:
    if isinstance(v, datetime.datetime) and v.year == 1905 and v.month <= 7:
        # a year (e.g. 2021) typed into a date-formatted cell shows as 1905-07-13
        year = (v.date() - datetime.date(1899, 12, 30)).days
        flags.append(f"{name}: only the year {year} given")
        return f"{year}-01-01"
    return _sane_date(to_date(v), name, v, flags)


def _is_million_col(f: dict, header: str) -> bool:
    return f["type"] == "decimal" and ("million" in header.lower() or "pkr million" in f["label"].lower())


def _rupee_columns(ws, hit: TableHit, colmap: dict, fields: dict, letters: list) -> set:
    """'PKR Million' columns that hold rupees: at least half the amounts are 10,000+ and hardly any (<= 10%) look
    like millions (< 1,000). Workbooks often keep the million header while typing rupees, so 1,000-9,999 there is
    rupees too (an Rs 8,000 XRD test, not 8 billion)."""
    out = set()
    for col, target in colmap.items():
        if target not in fields or not _is_million_col(fields[target], hit.headers.get(col, "")):
            continue
        idx = letters.index(col)
        nums = [n for r in range(hit.header_row + 1, hit.end_row + 1)
                if (n := to_number(ws.cell(r, idx + 1).value)) is not None and n > 0]
        if (len(nums) >= 3 and sum(n >= 10_000 for n in nums) >= 0.5 * len(nums)
                and sum(n < 1000 for n in nums) <= 0.1 * len(nums)):
            out.add(col)
    return out


def _import_table(batch, ws, hit: TableHit, spec: dict, matcher, res: ImportResult, user):
    cat = CATEGORY_BY_KEY[spec["category"]]
    colmap = _column_map(spec["columns"], hit.headers)
    fields = {f["name"]: f for f in cat["fields"]}
    scope = spec.get("scope")
    title_column = cat["title_field"] in colmap.values() or "_thematic_title" in colmap.values()
    count = flagged = 0
    letters = [get_column_letter(c) for c in range(1, ws.max_column + 1)]
    rupee_cols = _rupee_columns(ws, hit, colmap, fields, letters)
    for r in range(hit.header_row + 1, hit.end_row + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if spec.get("subheadings"):
            sh = _subheading(values)
            if sh:
                scope = sh
                continue
        if _is_noise(values):
            continue
        if spec.get("require_sn") and values[0] in (None, ""):
            res.skipped += 1
            continue
        source = {}
        for col, v in zip(letters, values):
            if v in (None, ""):
                continue
            head = hit.headers.get(col) or f"col {col}"
            source[head] = text(v) if not isinstance(v, (int, float)) else v
        data: dict = dict(spec.get("fixed") or {})
        flags: list[str] = []
        people_raw = {"lead": None, "co": None, "dept": None, "desig": None}
        evidence = []
        rupee_amount = None
        row_scope = scope
        for col, target in colmap.items():
            idx = letters.index(col)
            v = values[idx]
            if target is None or is_null(v):
                continue
            if target in ("_pi", "_lead"):
                people_raw["lead"] = v
            elif target == "_copi":
                people_raw["co"] = v
            elif target == "_dept":
                people_raw["dept"] = text(v)
            elif target == "_desig":
                people_raw["desig"] = text(v)
            elif target == "_duration":
                s, e = date_range(v)
                s, e = _sane_date(s, "start_date", v, flags), _sane_date(e, "end_date", v, flags)
                if s:
                    data.setdefault("start_date", s)
                if e:
                    data.setdefault("end_date", e)
            elif target == "_date_duration":
                d = _sane_date(to_date(v), cat["date_field"], v, flags)
                if d:
                    data[cat["date_field"]] = d
                data["agreement_duration"] = text(v)
            elif target == "_evidence":
                evidence += re.findall(r"https?://\S+", str(v)) or ([text(v)] if text(v) else [])
            elif target == "_amount_rupees":
                rupee_amount = rupees(v)
            elif target == "_thematic_title":
                data["research_proposal_title"] = text(v)
            elif target == "_scope":
                row_scope = nat_int(v) or row_scope
            elif target == "_scope_or_date":
                row_scope = nat_int(v) or row_scope
                if not nat_int(v) and to_date(v):
                    data.setdefault("approval_date", _sane_date(to_date(v), "approval_date", v, flags))
            elif target in fields:
                if fields[target]["type"] == "date":
                    data[target] = _clean_date(v, target, flags)
                else:
                    data[target] = _clean(fields[target], v, header=hit.headers.get(col, ""))
                    n = to_number(v) if _is_million_col(fields[target], hit.headers.get(col, "")) else None
                    if n is not None and 1000 <= n < 10_000:
                        if col in rupee_cols and float(n).is_integer():
                            data[target] = round(n / 1_000_000, 4)
                            flags.append(f"{target}: {text(v)} read as rupees (the column is mostly in rupees)")
                        else:
                            flags.append(f"{target}: {text(v)} read as PKR million; check the unit")
        if spec.get("fixed_fn"):
            data.update(spec["fixed_fn"](data))
        if spec.get("category_fn"):
            cat = CATEGORY_BY_KEY[spec["category_fn"](data)]
        slot = lead_slot(cat)
        co_slot = next((p for p in cat["people"] if not p["lead"]), None)
        if not row_scope and spec.get("scope_from") and data.get(spec["scope_from"]):
            authority = str(data[spec["scope_from"]])
            row_scope = nat_int(authority) or ("National" if re.search(r"\bipo\b", authority, re.I) else None)
        if row_scope and "national_or_international" in fields:
            data["national_or_international"] = row_scope
        if cat["amount_field"] and rupee_amount is not None:
            if cat["amount_unit"] == "M":
                data[cat["amount_field"]] = round(rupee_amount / 1_000_000, 4)
            else:
                data[cat["amount_field"]] = rupee_amount
        if cat["date_field"] and not data.get(cat["date_field"]):
            data[cat["date_field"]] = data.get("start_date")
        data = {k: v for k, v in data.items() if v not in (None, "")}
        meaningful = [k for k in data if k not in (spec.get("fixed") or {}) and k != "national_or_international"]
        if not meaningful and not people_raw["lead"]:
            res.skipped += 1
            continue

        # people
        contributors, owner, unmatched_text, suggestions = [], None, "", []
        leads = split_people(people_raw["lead"]) if people_raw["lead"] else []
        cos = split_people(people_raw["co"]) if people_raw["co"] else []
        if leads and (people_raw["dept"] or people_raw["desig"]):  # FY 2019-20 keeps them in their own columns
            leads[0]["dept_hint"] = " ".join(filter(None, [leads[0]["dept_hint"], people_raw["dept"]]))
            leads[0]["designation"] = leads[0]["designation"] or people_raw["desig"] or ""
        if slot and leads:
            first, rest = leads[0], leads[1:]
            person, sugg = matcher.match(first["name"], first["dept_hint"])
            if person:
                owner = person
                contributors.append(Contributor(role=slot["role"], person=person, order=0))
            else:
                unmatched_text, suggestions = first["raw"], sugg
                contributors.append(Contributor(role=slot["role"], name=first["name"][:255],
                                                designation=first["designation"][:120],
                                                affiliation=first["dept_hint"][:255], order=0))
            cos = rest + cos
        if co_slot:
            for i, c in enumerate(cos, start=1):
                person, _ = matcher.match(c["name"], c["dept_hint"])
                contributors.append(Contributor(role=co_slot["role"], person=person,
                                                name="" if person else c["name"][:255],
                                                designation="" if person else c["designation"][:120],
                                                affiliation="" if person else c["dept_hint"][:255], order=i))
        needs_match = bool(slot and leads and owner is None)

        title_field = cat["title_field"]
        if title_field and not data.get(title_field):
            lead_name = leads[0]["name"] if leads else ""
            prior = None
            if not title_column and cat["key"] in IP_FAMILY:
                prior = _recover_title(owner, lead_name, data, batch.cycle)
            if prior:
                data[title_field], data["_title_from"] = prior
                flags.append(f"title not in workbook; taken from earlier record #{prior[1]}")
            elif not title_column:
                described = next((data[k] for k in ("key_scientific_aspects", "key_initiatives",
                                                    "scope_of_collaboration", "field") if data.get(k)), None)
                data[title_field] = (described or f"{cat['short']} · {lead_name or 'untitled'}")[:255]
                flags.append("title not in workbook; placeholder used")
            else:
                # tables without a title value: fall back to the first text value
                data[title_field] = next((text(v) for v in values[1:] if isinstance(v, str) and len(v) > 3), None)
                if cat["key"] == "advisory_board_meeting" and data[title_field]:
                    # the title is the actions column; a blank one would otherwise show the meeting-date text alone
                    data[title_field] = f"{cat['short']} · {data[title_field]}"[:255]
                    flags.append("actions suggested not in workbook; placeholder used")
        if flags:
            data["_flags"] = flags
        data["_source"] = source
        if evidence:
            data["evidence_link"] = evidence[0]

        sub = Submission(category=cat["key"], cycle=batch.cycle, owner=owner, submitted_by=user, data=data,
                         status=Submission.Status.VERIFIED, source=Submission.Source.EXCEL, import_batch=batch,
                         import_ref=f"{hit.sheet}!{r}", needs_match=needs_match, verified_by=user,
                         verified_at=timezone.now(), submitted_at=timezone.now())
        if owner:
            sub.department = owner.department
            sub.faculty = owner.department.faculty if owner.department_id else None
        sub.promote()
        if not sub.event_date or not (batch.cycle.starts_on <= sub.event_date <= batch.cycle.ends_on):
            pass  # keep: the sheet says it belongs to this FY even when the typed date is off
        org_field = next((f["name"] for f in cat["fields"] if f["type"] == "org" and data.get(f["name"])), None)
        if org_field:
            sub.organization = resolve_organization(str(data[org_field])[:255])
        sub.save()
        for c in contributors:
            c.submission = sub
        Contributor.objects.bulk_create(contributors)
        sub.sync_departments()
        ReviewEvent.objects.create(submission=sub, actor=user, action=ReviewEvent.Action.IMPORT,
                                   to_status=sub.status, comment=f"{batch.original_name} {sub.import_ref}")
        ImportRow.objects.create(
            batch=batch, ref=sub.import_ref, table=hit.key, category=cat["key"],
            status=ImportRow.Status.UNMATCHED if needs_match else ImportRow.Status.CREATED,
            raw=source, parsed={k: v for k, v in data.items() if k != "_source"}, person_text=unmatched_text[:500],
            suggestions=suggestions, submission=sub, message="; ".join(flags))
        count += 1
        flagged += int(bool(flags))
        res.created += 1
        res.unmatched += int(needs_match)
    res.per_table[hit.key] = {"category": spec["category"], "rows": count, "sheet": hit.sheet,
                              "title": hit.title[:120], "flagged": flagged,
                              "mapped_columns": {hit.headers.get(c) or c: t for c, t in colmap.items()}}
    if flagged:
        res.flags[hit.key] = flagged
    res.log.append(f"{hit.key} ({hit.sheet} rows {hit.header_row + 1}-{hit.end_row}): {count} records"
                   + (f", {flagged} flagged" if flagged else ""))


def _recover_title(owner, lead_name: str, data: dict, cycle):
    """IP rows whose workbook lost the title column: reuse the title of the same lead's earlier IP record when
    exactly one earlier title agrees with at least one other detail (status, filing body, licensee …)."""
    qs = Submission.objects.filter(category__in=IP_FAMILY, cycle__year__lt=cycle.year).exclude(title="")
    if owner:
        qs = qs.filter(owner=owner)
    elif lead_name:
        qs = qs.filter(contributors__order=0, contributors__name__iexact=lead_name)
    else:
        return None
    keys = [k for k in ("development_status", "filed_with", "licensee", "commercial_partner",
                        "key_scientific_aspects", "field_of_use", "industrial_partner", "disclosure_made_with")
            if data.get(k)]
    best, found = 0, {}
    for s in qs.only("id", "title", "data"):
        score = sum(str(s.data.get(k) or "").strip().lower() == str(data[k]).strip().lower() for k in keys)
        if score > best:
            best, found = score, {s.title: s.pk}
        elif score == best and score:
            found.setdefault(s.title, s.pk)
    if best >= 1 and len(found) == 1:
        title, pk = next(iter(found.items()))
        return title, pk
    return None


def _import_aggregate(batch, ws, hit: TableHit, spec: dict, res: ImportResult, user):
    """One record for the whole table (FY 2019-20 incubation centre): each column's values joined by line."""
    cat = CATEGORY_BY_KEY[spec["category"]]
    colmap = _column_map(spec["columns"], hit.headers)
    letters = [get_column_letter(c) for c in range(1, ws.max_column + 1)]
    collected: dict[str, list] = {}
    source = []
    for r in range(hit.header_row + 1, hit.end_row + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if _is_noise(values):
            continue
        source.append({hit.headers.get(c) or f"col {c}": text(v) for c, v in zip(letters, values)
                       if v not in (None, "")})
        for col, target in colmap.items():
            v = text(values[letters.index(col)])
            if target and v and v not in collected.setdefault(target, []):
                collected[target].append(v)
    if not collected:
        res.per_table[hit.key] = {"category": cat["key"], "rows": 0, "sheet": hit.sheet}
        return
    data = {t: "\n".join(vals) for t, vals in collected.items()}
    if "status" in collected:
        data["status"] = collected["status"][0]
    data["_source"] = {"rows": source}
    sub = Submission(category=cat["key"], cycle=batch.cycle, submitted_by=user, data=data,
                     status=Submission.Status.VERIFIED, source=Submission.Source.EXCEL, import_batch=batch,
                     import_ref=f"{hit.sheet}!{hit.header_row + 1}-{hit.end_row}", verified_by=user,
                     verified_at=timezone.now(), submitted_at=timezone.now())
    sub.promote()
    sub.save()
    ReviewEvent.objects.create(submission=sub, actor=user, action=ReviewEvent.Action.IMPORT, to_status=sub.status,
                               comment=f"{batch.original_name} {sub.import_ref}")
    ImportRow.objects.create(batch=batch, ref=sub.import_ref, table=hit.key, category=cat["key"],
                             status=ImportRow.Status.CREATED, raw={"rows": source},
                             parsed={k: v for k, v in data.items() if k != "_source"}, submission=sub)
    res.created += 1
    res.per_table[hit.key] = {"category": cat["key"], "rows": 1, "sheet": hit.sheet, "title": hit.title[:120],
                              "source_rows": len(source)}
    res.log.append(f"{hit.key} ({hit.sheet}): 1 record from {len(source)} rows")


def _clean(f: dict, v, *, header: str = ""):
    t = f["type"]
    if t == "date":
        return to_date(v)
    if t == "decimal":
        if "million" in header.lower() or "(pkr million" in f["label"].lower() or "pkr million" in f["label"].lower():
            return million(v)
        return rupees(v) if "pkr" in f["label"].lower() else to_number(v)
    if t == "int":
        n = to_number(v)
        return int(n) if n is not None else None
    if t in ("select", "radio"):
        s = text(v)
        if not s:
            return None
        for opt in f.get("options") or []:
            if s.lower() == opt.lower() or s.lower().startswith(opt.lower()):
                return opt
        if f["name"] == "national_or_international":
            return nat_int(s)
        return s
    return text(v)


# ---------------------------------------------------------------------------
# Index-A (Annexure 1 in FY 2019-20) / Score Card / strategic plan
# ---------------------------------------------------------------------------
def _import_index_a(wb, cycle, user, res, *, template="index_2021", as_submitted=None):
    from scorecard.models import InstitutionProfile

    sheet = "Annexure 1" if template == "annexure_2019" else "Index-A"
    if sheet not in wb:
        return
    ws = wb[sheet]
    head_row = next((r for r in range(1, 20) if str(ws.cell(r, 1).value or "").strip().upper() == "SN"
                     or str(ws.cell(r, 2).value or "").strip().lower() == "position"), 6)
    details = {}
    for r in range(1, head_row):
        for c in range(2, ws.max_column, 2):  # label in B, D, F … value right next to it
            k, v = ws.cell(r, c).value, ws.cell(r, c + 1).value
            if isinstance(k, str) and k.strip() and not is_null(v):
                details[" ".join(k.split()).rstrip(" :")] = v if isinstance(v, (int, float)) else text(v)
    heads = {c: " ".join(str(ws.cell(head_row, c).value or "").split()).lower() for c in range(1, 16)}

    def col(word):
        return next((c for c, h in heads.items() if word in h), None)

    staff = []
    for r in range(head_row + 1, head_row + 30):
        name = ws.cell(r, col("name") or 4).value
        pos = ws.cell(r, col("position") or 2).value
        if is_null(name) and is_null(pos):
            if staff:
                break
            continue
        staff.append({
            "position": text(pos), "employment": text(ws.cell(r, col("full time") or 3).value), "name": text(name),
            "office_phone": text(ws.cell(r, col("office") or 5).value),
            "mobile": text(ws.cell(r, col("mobile") or 6).value), "email": text(ws.cell(r, col("email") or 7).value),
            "qualification_level": text(ws.cell(r, col("qualification level") or 8).value),
            "qualification_title": text(ws.cell(r, col("qualification title") or 9).value),
            "field_of_study": text(ws.cell(r, col("field of study") or 10).value),
            "appointment_date": to_date(ws.cell(r, col("date of appointment") or 12).value),
            "experience": text(ws.cell(r, col("total experience") or 14).value),
        })
    profile, _ = InstitutionProfile.objects.get_or_create(cycle=cycle)
    profile.details = details
    profile.staff = staff
    ops = _operations_from_score_card(wb) if template == "index_2021" else _operations_from_card(as_submitted)
    profile.operations = {**(profile.operations or {}), **ops}
    profile.updated_by = user
    profile.save()
    for label, attr in ((r"phd faculty", "phd_faculty_count"), (r"sanctioned", "sanctioned_faculty_posts"),
                        (r"existing faculty", "existing_faculty_count")):
        v = next((v for k, v in details.items() if re.search(label, k.lower())), None)
        if isinstance(v, (int, float)):
            setattr(cycle, attr, int(v))
    cycle.save()
    res.log.append(f"{sheet}: {len(details)} profile fields, {len(staff)} ORIC staff, PhD faculty = "
                   f"{cycle.phd_faculty_count}")


def _operations_from_score_card(wb) -> dict:
    if "Score Card" not in wb.sheetnames:
        return {}
    ws = wb["Score Card"]
    keys = {"i": "head", "ii": "manager_rm", "iii": "manager_ic", "iv": "support", "v": "steering_committee",
            "vi": "sc_meetings", "vii": "irb_meetings", "viii": "offices_ict", "ix": "webpage", "x": "database"}
    out = {}
    for r in range(5, 15):
        item = str(ws.cell(r, 2).value or "").strip().lower()
        v = ws.cell(r, 5).value
        if item in keys:
            k = keys[item]
            if k in ("sc_meetings", "irb_meetings"):
                out[k] = to_number(v) or 0
            elif k == "database":
                out[k] = "excel" if str(v).lower().startswith("y") else "no"
            else:
                out[k] = str(v).strip().lower().startswith("y")
    return out


# Section A items of older score cards, recognised by wording; "present" = a self-score above zero.
OPS_WORDING = [
    ("sc_meetings", None), ("irb_meetings", None),
    ("steering_committee", r"steering committee"), ("head", r"oric head|director|pro-?rector"),
    ("manager_rm", r"manager research"), ("manager_ic", r"manager (innovation|intellectual)"),
    ("support", r"support position|administrative|accountant"), ("webpage", r"webpage"),
    ("database", r"database"), ("offices_ict", r"office|ict facilit|available facilit"),
]


def _operations_from_card(card: dict | None) -> dict:
    if not card or not card.get("sections"):
        return {}
    out = {}
    for row in card["sections"][0]["rows"]:
        words = f"{row.get('kpi') or ''} {row.get('parameter') or ''}".lower()
        present = isinstance(row.get("self"), (int, float)) and row["self"] > 0
        for key, pattern in OPS_WORDING:
            if pattern and key not in out and re.search(pattern, words):
                out[key] = ("excel" if present else "no") if key == "database" else present
                break
    return out


def _sheet_scores(wb, res):
    """Keep the workbook's own reported numbers / self-scores for reconciliation in the UI."""
    if "Score Card" not in wb.sheetnames:
        return
    ws = wb["Score Card"]
    section = None
    out = {}
    for r in range(4, 66):
        a = str(ws.cell(r, 1).value or "").strip()
        if re.match(r"^[ABCD]\.?$", a):
            section = a[0]
            continue
        item = str(ws.cell(r, 2).value or "").strip()
        if not item or not section:
            continue
        reported = ws.cell(r, 6).value if section != "A" else ws.cell(r, 5).value
        if reported is None:
            reported = ws.cell(r, 5).value
        out[f"{section}.{item}"] = {"reported": reported, "max": ws.cell(r, 7).value, "score": ws.cell(r, 8).value}
    res.per_table["_sheet_scorecard"] = out


ROMAN_RE = re.compile(r"^(?:[ivx]+[a-c]?|[ivx]+-[a-c])$", re.I)


def _num(v):
    if isinstance(v, (int, float)):
        return round(float(v), 4)
    return text(v)


def score_card_as_submitted(wb, template: str) -> dict:
    """The Score Card exactly as filed with HEC, for any template: sections → items with the reported number,
    max score, self-assessed score and HEC's score. Columns are found by their header words."""
    if "Score Card" not in wb:
        return {}
    ws = wb["Score Card"]
    head = next((r for r in range(1, 15) for c in range(1, 6)
                 if isinstance(ws.cell(r, c).value, str) and "key performance" in ws.cell(r, c).value.lower()), None)
    if not head:
        return {}
    cols: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        h = " ".join(str(ws.cell(head, c).value or "").split()).lower()
        for key, pattern in (("kpi", r"key performance"), ("parameter", r"parameters"), ("max", r"^max"),
                             ("self", r"self"), ("hec", r"hec"), ("reported", r"reported"),
                             ("evidence_page", r"evidence at page")):
            if key not in cols and re.search(pattern, h):
                cols[key] = c
                break
    kpi_col = cols.get("kpi", 2)
    item_col = kpi_col - 1 if kpi_col > 1 else None
    # 2020+ cards put area labels in column A and the roman item in B; 2019-20 has the item in A.
    if item_col and item_col > 1 and not any(ROMAN_RE.match(str(ws.cell(r, item_col).value or "").strip())
                                             for r in range(head + 1, min(ws.max_row, head + 30) + 1)):
        item_col -= 1

    def val(r, key):
        return _num(ws.cell(r, cols[key]).value) if key in cols else None

    sections, cur, kpi, area = [], None, None, None
    total = None
    for r in range(head + 1, ws.max_row + 1):
        cells = [text(ws.cell(r, c).value) or "" for c in range(1, min(ws.max_column, 12) + 1)]
        joined = " ".join(c for c in cells if c)
        if not joined:
            continue
        low = joined.lower()
        if re.search(r"ranking criteria|^\*|^note", low):
            break
        a = cells[0]
        m = re.match(r"^([A-D])\.(?:\s|$)|^(\d)\.\s", a)
        if m and val(r, "max") in (None, "") and val(r, "self") in (None, ""):
            title = next((c for c in cells[1:] if len(c) > 3), "") or re.sub(r"^\S+\s*", "", a)
            cur = {"code": m.group(1) or m.group(2), "title": title, "rows": [], "subtotal": None}
            sections.append(cur)
            kpi = area = None
            continue
        if re.search(r"grand total|^total\b|total score", low) and "sub total" not in low:
            total = {"max": val(r, "max"), "self": val(r, "self"), "hec": val(r, "hec")}
            continue
        if cur is None:
            continue
        if "sub total" in low:
            cur["subtotal"] = {"max": val(r, "max"), "self": val(r, "self"), "hec": val(r, "hec")}
            continue
        item = str(ws.cell(r, item_col).value or "").strip() if item_col else ""
        if not ROMAN_RE.match(item):
            if item_col == 1 or not item:
                item = ""
            else:  # an area label in column A ("Human Resource & Management")
                area, item = item, ""
        if item_col and item_col > 1 and cells[0] and not m:
            area = cells[0]
        kpi = text(ws.cell(r, kpi_col).value) or kpi
        mx, sf = val(r, "max"), val(r, "self")
        reported = val(r, "reported")
        if reported in (None, "") and isinstance(val(r, "evidence_page"), float):
            reported = val(r, "evidence_page")  # FY 2024-25 card shifted section B-D numbers one column right
        if mx in (None, "") and sf in (None, "") and reported in (None, ""):
            continue
        cur["rows"].append({"item": item, "area": area, "kpi": kpi, "parameter": val(r, "parameter"),
                            "reported": reported, "max": mx, "self": sf, "hec": val(r, "hec")})
    if total is None and sections:
        def tot(k):
            vals = [s["subtotal"][k] for s in sections if s["subtotal"] and isinstance(s["subtotal"][k], float)]
            return round(sum(vals), 4) if vals else None
        total = {"max": tot("max"), "self": tot("self"), "hec": tot("hec")}
    reported_head = text(ws.cell(head, cols["reported"]).value) if "reported" in cols else None
    fy = re.search(r"FY\s*(\d{2,4})\s*-\s*(\d{2,4})", reported_head or "")
    reported_fy = f"FY {int(fy.group(1)) % 100 + 2000}-{int(fy.group(2)) % 100:02d}" if fy else None
    return {"template": template, "framework": TEMPLATE_FRAMEWORK[template], "reported_fy": reported_fy,
            "sections": sections, "total": total}


def _import_strategic_plan(batch, wb, cycle, user, res):
    if "Index-D" not in wb.sheetnames:
        return
    ws = wb["Index-D"]
    for r in range(1, 6):
        if str(ws.cell(r, 1).value or "").strip().lower() == "ia":
            title = text(ws.cell(r, 2).value) or "Five-Year Strategic Plan"
            link = next((str(ws.cell(r, c).value) for c in range(3, 8)
                         if str(ws.cell(r, c).value or "").startswith("http")), None)
            sub = Submission(category="strategic_plan", cycle=cycle, submitted_by=user, source=Submission.Source.EXCEL,
                             import_batch=batch, import_ref=f"Index-D!{r}", status=Submission.Status.VERIFIED,
                             data={"title": title, **({"evidence_link": link} if link else {}),
                                   "_source": {"row": [text(ws.cell(r, c).value) for c in range(1, 6)]}},
                             verified_at=timezone.now(), verified_by=user)
            sub.promote()
            sub.save()
            res.created += 1
            res.per_table["D.ia"] = {"category": "strategic_plan", "rows": 1}
            return
