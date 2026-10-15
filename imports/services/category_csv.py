"""Bulk import of one category from CSV / XLSX (our template, or a Google Form response export).

Columns are matched to registry fields by field name or label (case/punctuation-insensitive,
fuzzy for Google Form question text). Person columns: ``lead_email`` / ``email`` / "Email address of
Faculty" pick the owner; ``lead_name`` / "Name of Faculty Member" / "Name of PI" are matched by name.
"""
from __future__ import annotations

import csv
import io
import re
from difflib import SequenceMatcher

from django.db import transaction
from django.utils import timezone

from people.models import Person
from submissions.models import Contributor, Submission
from submissions.registry import CATEGORY_BY_KEY, lead_slot

from ..models import ImportBatch, ImportRow
from .hec_workbook import _clean
from .normalize import PersonMatcher, split_people, text


def _key(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (s or "").lower()).strip()


def template_csv(category: str) -> str:
    c = CATEGORY_BY_KEY[category]
    cols = ["fiscal_year", "lead_email", "lead_name", "co_contributors"] + \
        [f["name"] for f in c["fields"] if f["type"] != "file"]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(cols)
    w.writerow(["FY 2025-26", "someone@neduet.edu.pk", "", "Dr. A (external); colleague@neduet.edu.pk"] +
               [_example(f) for f in c["fields"] if f["type"] != "file"])
    return buf.getvalue()


def _example(f):
    if f.get("options"):
        return f["options"][0]
    return {"date": "2025-09-01", "decimal": "1.5", "int": "10", "url": "https://..."}.get(f["type"], "")


def read_table(name: str, content: bytes) -> list[dict]:
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        head = [str(h or "").strip() for h in rows[0]]
        return [dict(zip(head, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
    textv = content.decode("utf-8-sig", errors="replace")
    return list(csv.DictReader(io.StringIO(textv)))


def map_columns(category: str, headers: list[str]) -> dict:
    c = CATEGORY_BY_KEY[category]
    fields = [f for f in c["fields"] if f["type"] != "file"]
    out = {}
    for h in headers:
        k = _key(h)
        if k in ("lead email", "email", "email address of faculty", "email address of faculty member",
                 "email address"):
            out[h] = "_lead_email"
            continue
        if k in ("lead name", "name of faculty member", "full name of faculty member and designation") or \
                k.startswith("name of pi"):
            out[h] = "_lead_name"
            continue
        if k in ("co contributors",) or k.startswith("name of co pi"):
            out[h] = "_co"
            continue
        if k in ("fiscal year", "fy"):
            out[h] = "_fy"
            continue
        best, score = None, 0.0
        for f in fields:
            for cand in (f["name"], f["label"]):
                s = 1.0 if _key(cand) == k else SequenceMatcher(None, _key(cand), k).ratio()
                if s > score:
                    best, score = f["name"], s
        if best and score >= 0.72 and best not in out.values():
            out[h] = best
    return out


class _Rollback(Exception):
    pass


def import_category(batch: ImportBatch, category: str, rows: list[dict], *, user, dry_run=False,
                    status=Submission.Status.SUBMITTED) -> dict:
    c = CATEGORY_BY_KEY[category]
    colmap = map_columns(category, list(rows[0].keys()) if rows else [])
    fields = {f["name"]: f for f in c["fields"]}
    slot = lead_slot(c)
    co_slot = next((p for p in c["people"] if not p["lead"]), None)
    matcher = PersonMatcher()
    stats = {"created": 0, "unmatched": 0, "errors": 0, "columns": colmap}
    try:
        with transaction.atomic():
            for i, raw in enumerate(rows, start=2):
                data, owner, unmatched, sugg, cos = {}, None, "", [], []
                for col, target in colmap.items():
                    v = raw.get(col)
                    if v in (None, ""):
                        continue
                    if target == "_lead_email":
                        owner = owner or Person.objects.filter(user__email__iexact=str(v).strip()).first()
                    elif target == "_lead_name" and not owner:
                        ppl = split_people(v)
                        if ppl:
                            owner, sugg = matcher.match(ppl[0]["name"], ppl[0]["dept_hint"])
                            if not owner:
                                unmatched = ppl[0]["raw"]
                    elif target == "_co":
                        for part in str(v).split(";"):
                            part = part.strip()
                            if not part:
                                continue
                            p = Person.objects.filter(user__email__iexact=part).first() if "@" in part else \
                                matcher.match(part)[0]
                            cos.append((p, part))
                    elif target in fields:
                        data[target] = _clean(fields[target], v)
                data = {k: v for k, v in data.items() if v not in (None, "")}
                if not data:
                    continue
                data["_source"] = {k: text(v) for k, v in raw.items() if v not in (None, "")}
                sub = Submission(category=category, cycle=batch.cycle, owner=owner, submitted_by=user, data=data,
                                 status=status, source=Submission.Source.CSV if batch.kind == ImportBatch.Kind.
                                 CATEGORY_CSV else Submission.Source.FORM, import_batch=batch,
                                 import_ref=f"row {i}", needs_match=bool(slot and not owner and unmatched),
                                 submitted_at=timezone.now())
                if owner and owner.department_id:
                    sub.department, sub.faculty = owner.department, owner.department.faculty
                sub.promote()
                sub.save()
                if slot:
                    Contributor.objects.create(submission=sub, role=slot["role"], person=owner,
                                               name="" if owner else unmatched[:255], order=0)
                for n, (p, txt) in enumerate(cos, start=1):
                    if co_slot:
                        Contributor.objects.create(submission=sub, role=co_slot["role"], person=p,
                                                   name="" if p else txt[:255], order=n)
                sub.sync_departments()
                ImportRow.objects.create(batch=batch, ref=f"row {i}", category=category,
                                         status=ImportRow.Status.UNMATCHED if sub.needs_match
                                         else ImportRow.Status.CREATED,
                                         raw=data["_source"], parsed={k: v for k, v in data.items() if k != "_source"},
                                         person_text=unmatched[:500], suggestions=sugg, submission=sub)
                stats["created"] += 1
                stats["unmatched"] += int(sub.needs_match)
            if dry_run:
                raise _Rollback
    except _Rollback:
        pass
    return stats
