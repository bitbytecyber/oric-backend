"""Submission exports: flat table (CSV / XLSX), one-sheet-per-category workbook, evidence ZIP."""
from __future__ import annotations

import csv
import io
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from submissions.models import Submission
from submissions.registry import CATEGORIES, CATEGORY_BY_KEY

HEADER_FILL = PatternFill("solid", fgColor="1F6F7A")
HEADER_FONT = Font(bold=True, color="FFFFFF")

# key -> (header, getter)
BASE_COLUMNS: dict[str, tuple[str, callable]] = {
    "id": ("ID", lambda s: s.pk),
    "fiscal_year": ("Fiscal year", lambda s: s.cycle.label),
    "pillar": ("Pillar", lambda s: s.get_pillar_display()),
    "category": ("Category", lambda s: s.category_def["short"]),
    "hec_ref": ("HEC ref", lambda s: s.category_def["excel"]),
    "form_ref": ("Form ref", lambda s: s.category_def["form"]),
    "title": ("Title", lambda s: s.title),
    "owner": ("PI / lead", lambda s: s.owner.full_name if s.owner_id else _lead_text(s)),
    "designation": ("Designation", lambda s: s.owner.designation if s.owner_id else ""),
    "department": ("Department", lambda s: s.department.code if s.department_id else ""),
    "faculty": ("Faculty", lambda s: s.faculty.code if s.faculty_id else ""),
    "contributors": ("Contributors", lambda s: "; ".join(
        f"{c.display_name} ({c.role})" for c in s.contributors.all())),
    "event_date": ("Date", lambda s: s.event_date),
    "amount_pkr": ("Amount (PKR)", lambda s: float(s.amount_pkr) if s.amount_pkr is not None else None),
    "scope_level": ("National / International", lambda s: s.scope_level),
    "funding_source": ("Funding source", lambda s: s.get_funding_source_display() if s.funding_source else ""),
    "organization": ("Organization", lambda s: s.organization.name if s.organization_id else ""),
    "status": ("Status", lambda s: s.get_status_display()),
    "submitted_at": ("Submitted at", lambda s: s.submitted_at.replace(tzinfo=None) if s.submitted_at else None),
    "verified_at": ("Verified at", lambda s: s.verified_at.replace(tzinfo=None) if s.verified_at else None),
    "entered_by": ("Entered by", lambda s: s.submitted_by.full_name if s.submitted_by_id else ""),
    "evidence": ("Evidence links", lambda s: " ".join(s.evidence_links or [])),
    "evidence_files": ("Evidence files", lambda s: "; ".join(f.original_name for f in s.files.all())),
    "source": ("Source", lambda s: s.get_source_display()),
    "import_ref": ("Import ref", lambda s: s.import_ref),
}
DEFAULT_COLUMNS = ["id", "fiscal_year", "category", "hec_ref", "title", "owner", "designation", "department",
                   "faculty", "contributors", "event_date", "amount_pkr", "scope_level", "funding_source",
                   "organization", "status", "evidence", "source"]


def _lead_text(s: Submission) -> str:
    c = next((c for c in s.contributors.all()), None)
    return c.display_name if c else ""


def prepare(qs):
    return qs.select_related("cycle", "owner", "department", "faculty", "organization", "submitted_by__person") \
        .prefetch_related("contributors__person", "files").order_by("category", "event_date", "pk")


def table_rows(qs, columns: list[str]):
    cols = [c for c in columns if c in BASE_COLUMNS] or DEFAULT_COLUMNS
    yield [BASE_COLUMNS[c][0] for c in cols]
    for s in qs:
        yield [BASE_COLUMNS[c][1](s) for c in cols]


def to_csv(qs, columns) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    for row in table_rows(qs, columns):
        w.writerow(["" if v is None else v for v in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _style_header(ws, ncols):
    for i in range(1, ncols + 1):
        cell = ws.cell(row=1, column=i)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _autosize(ws):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col[:200] if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, width + 2), 60)


def to_xlsx(qs, columns, *, title="Submissions") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    rows = list(table_rows(qs, columns))
    for r in rows:
        ws.append(r)
    _style_header(ws, len(rows[0]))
    _autosize(ws)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_xlsx_by_category(qs) -> bytes:
    """One sheet per category with every registry field as its own column."""
    wb = Workbook()
    wb.remove(wb.active)
    summary = wb.create_sheet("Summary")
    summary.append(["Category", "HEC ref", "Form ref", "Records", "Amount (PKR)"])
    qs = list(qs)
    by_cat: dict[str, list] = {}
    for s in qs:
        by_cat.setdefault(s.category, []).append(s)
    for c in CATEGORIES:
        items = by_cat.get(c["key"])
        if not items:
            continue
        fields = [f for f in c["fields"] if f["type"] != "file"]
        ws = wb.create_sheet(c["short"][:31].replace("/", "-"))
        head = ["ID", "Fiscal year", "Status", "PI / lead", "Department", "Contributors",
                *[f["label"] for f in fields], "Evidence files", "Evidence links"]
        ws.append(head)
        for s in items:
            ws.append([s.pk, s.cycle.label, s.get_status_display(),
                       s.owner.full_name if s.owner_id else _lead_text(s),
                       s.department.code if s.department_id else "",
                       "; ".join(f"{x.display_name} ({x.role})" for x in s.contributors.all()),
                       *[s.data.get(f["name"]) for f in fields],
                       "; ".join(f.original_name for f in s.files.all()), " ".join(s.evidence_links or [])])
        _style_header(ws, len(head))
        _autosize(ws)
        summary.append([c["short"], c["excel"], c["form"], len(items),
                        float(sum((s.amount_pkr or 0) for s in items))])
    _style_header(summary, 5)
    _autosize(summary)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def evidence_zip(qs) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        index = ["submission_id,category,title,file"]
        for s in qs:
            for f in s.files.all():
                name = f"{CATEGORY_BY_KEY[s.category]['short'][:40]}/{s.pk}-{f.field}-{f.original_name}"
                try:
                    with f.file.open("rb") as fh:
                        z.writestr(name, fh.read())
                    index.append(f'{s.pk},{s.category},"{s.title.replace(chr(34), "")}",{name}')
                except FileNotFoundError:
                    continue
        z.writestr("index.csv", "\n".join(index))
    return buf.getvalue()
