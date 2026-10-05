import csv
import io

from django.db.models import Count, Q
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from audit.services import record_export
from org.models import ReportingCycle
from rbac.utils import user_has_capability
from rbac.viewsets import CapabilityAPIViewMixin
from submissions.access import scope_q
from submissions.models import SectionEntry, Submission
from submissions.registry import PILLAR_BY_KEY, PILLARS

from . import services


def _int(v):
    try:
        return int(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


class ScorecardSummaryView(CapabilityAPIViewMixin, APIView):
    required_capability = "report.view"

    def get(self, request):
        years = services.available_years()
        year = _int(request.query_params.get("year")) or (years[0] if years else None)
        data = services.scorecard_summary(
            request.user, year=year, department=_int(request.query_params.get("department")),
            pillar=request.query_params.get("pillar") or None,
            approved_only=request.query_params.get("approved_only") in ("1", "true"),
        )
        data["years"] = years
        data["trend"] = services.yearly_trend(request.user, pillar=request.query_params.get("pillar") or None)
        return Response(data)


class HomeView(APIView):
    """Personal home dashboard: my returns this cycle (+ institution KPIs for dashboard.view_admin)."""

    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        cycle = ReportingCycle.current()
        mine = {}
        if cycle:
            for s in Submission.objects.filter(owner=user, cycle=cycle).annotate(n=Count("entries")):
                mine[s.pillar] = {"id": s.pk, "status": s.status, "status_label": s.get_status_display(),
                                  "entries": s.n, "updated_at": s.updated_at,
                                  "missing": len(s.missing_entries())}
        out = {
            "cycle": {"id": cycle.pk, "label": cycle.label, "year": cycle.year, "status": cycle.status,
                      "deadline": cycle.submission_deadline, "accepts_submissions": cycle.accepts_submissions}
            if cycle else None,
            "my_returns": [
                {"pillar": p["key"], "name": p["name"], "tagline": p["tagline"], "return": mine.get(p["key"])}
                for p in PILLARS
            ],
            "admin": None,
        }
        if cycle and user_has_capability(user, "dashboard.view_admin"):
            subs = Submission.objects.filter(cycle=cycle)
            q = scope_q(user, "submission.view")
            if q is not None:
                subs = subs.filter(q)
            review_q = scope_q(user, "submission.review")
            to_review = Submission.objects.filter(cycle=cycle, status__in=["submitted", "under_review"])
            if review_q is not None:
                to_review = to_review.filter(review_q)
            by_status = dict(subs.values("status").annotate(n=Count("id")).values_list("status", "n"))
            out["admin"] = {
                "submissions": subs.count(),
                "faculty": subs.values("owner").distinct().count(),
                "departments": subs.values("department").distinct().count(),
                "by_status": by_status,
                "awaiting_review": to_review.count() if user_has_capability(user, "submission.review") else None,
                "by_pillar": [
                    {"pillar": p["key"], "name": p["name"],
                     "total": subs.filter(pillar=p["key"]).count(),
                     "approved": subs.filter(pillar=p["key"], status="approved").count()}
                    for p in PILLARS
                ],
                "recent": [
                    {"id": s.pk, "pillar": s.pillar, "faculty_name": s.faculty_name, "status": s.status,
                     "department": s.department.name, "submitted_at": s.submitted_at}
                    for s in subs.exclude(status="draft").select_related("department").order_by("-submitted_at")[:8]
                ],
            }
        return Response(out)


class ExportView(CapabilityAPIViewMixin, APIView):
    """CSV (headers only) or XLSX (headers + one sheet per section) of returns in scope."""

    required_capability = "submission.export"
    envelope = False

    def get(self, request):
        pillar = request.query_params.get("pillar")
        if pillar not in PILLAR_BY_KEY:
            return HttpResponse("pillar must be RE, IC or SCB", status=400)
        fmt = request.query_params.get("format", "xlsx")
        qs = Submission.objects.filter(pillar=pillar).select_related("department", "cycle")
        q = scope_q(request.user, "submission.export")
        if q is not None:
            qs = qs.filter(q)
        if _int(request.query_params.get("year")):
            qs = qs.filter(cycle__year=_int(request.query_params.get("year")))
        if _int(request.query_params.get("department")):
            qs = qs.filter(department_id=_int(request.query_params.get("department")))
        if request.query_params.get("status"):
            qs = qs.filter(status__in=request.query_params.get("status").split(","))
        p = PILLAR_BY_KEY[pillar]
        head = ["Return ID", "Year", "Faculty name", "Email", "Designation", "Department", "Status",
                "Submitted at"] + [c["label"] for c in p["counts"]]

        def head_row(s):
            return [s.pk, s.cycle.year, s.faculty_name, s.faculty_email, s.designation, s.department.name,
                    s.get_status_display(), s.submitted_at.strftime("%Y-%m-%d %H:%M") if s.submitted_at else ""] + [
                (s.counts or {}).get(c["name"], 0) for c in p["counts"]]

        subs = list(qs.order_by("department__name", "faculty_name"))
        record_export(request, description=f"Export {pillar} returns ({fmt})", count=len(subs))
        filename = f"oric-{pillar.lower()}-{request.query_params.get('year') or 'all'}"

        if fmt == "csv":
            buf = io.StringIO()
            w = csv.writer(buf)
            w.writerow(head)
            for s in subs:
                w.writerow(head_row(s))
            resp = HttpResponse(buf.getvalue(), content_type="text/csv")
            resp["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
            return resp

        wb = Workbook()
        ws = wb.active
        ws.title = "Returns"
        bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F6F7A")
        ws.append(head)
        for cell in ws[1]:
            cell.font, cell.fill = bold, fill
        for s in subs:
            ws.append(head_row(s))
        by_sub = {s.pk: s for s in subs}
        for sec in p["sections"]:
            sh = wb.create_sheet(sec["key"])
            fields = [f for f in sec["fields"] if f["type"] != "file"]
            sh.append(["Return ID", "Faculty name", "Department"] + [f["label"] for f in fields])
            for cell in sh[1]:
                cell.font, cell.fill = bold, fill
            for e in SectionEntry.objects.filter(section=sec["key"], submission_id__in=by_sub.keys()):
                s = by_sub[e.submission_id]
                sh.append([s.pk, s.faculty_name, s.department.name] + [e.data.get(f["name"]) for f in fields])
        out = io.BytesIO()
        wb.save(out)
        resp = HttpResponse(out.getvalue(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
        return resp
